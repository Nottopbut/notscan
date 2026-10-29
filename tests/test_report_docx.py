"""DOCX-отчёт: landscape, фиксированные читаемые колонки, переносы, шапка.

Регресс на жалобу «странно строится отчёт в docx: колонки узкие, текст
нечитаемый»: раньше таблица Findings из 7 колонок авто-сжималась на
портретной A4 без фиксированных ширин.
"""

import os
import tempfile

import pytest

import db
import report_export


@pytest.fixture(autouse=True)
def _english_reports():
    """Экспорт отчётов локализуется по языку UI; тест проверяет EN-вывод."""
    import i18n
    saved = i18n.get_language()
    i18n.set_language("en")
    try:
        yield
    finally:
        i18n.set_language(saved)


def _seed(tmp_db: str) -> int:
    db.set_db_path(tmp_db)
    db.init_db()
    rows = [
        ("http://target.test/search?q=" + "x" * 120, "sqli", "SQL error: you have an error in your sql syntax " * 4),
        ("http://target.test/user?id=1234567890", "idor", "Object-level authorization missing"),
        ("http://target.test/hostheader", "host_header_injection", "Host reflected in body/absolute URL " * 3),
    ]
    for url, ct, ev in rows:
        rid = db.insert_scan_result(
            template_id=f"t-{ct}", worker_id=0, method="GET", url=url,
            status_code=200, error=None, check_type=ct, payload="p", evidence=ev,
        )
        db.upsert_scan_meta(scan_result_id=rid, detector="matcher",
                            severity="HIGH", confidence=0.7, exploitability=0.6)
    return len(rows)


def test_docx_landscape_and_readable_columns():
    from docx import Document
    from docx.enum.section import WD_ORIENT

    tmp_db = tempfile.mktemp(suffix=".db")
    out = tempfile.mktemp(suffix=".docx")
    try:
        assert _seed(tmp_db) == 3
        n = report_export.export_docx(out)
        assert n == 3
        doc = Document(out)
        sec = doc.sections[0]
        assert sec.orientation == WD_ORIENT.LANDSCAPE
        assert sec.page_width > sec.page_height
        assert len(doc.tables) >= 4

        # таблица Findings — 7 колонок, ширины фиксированы и не «огрызки»
        findings = doc.tables[-1]
        assert len(findings.columns) == 7
        widths = [
            (findings.rows[0].cells[i].width.inches if findings.rows[0].cells[i].width else 0)
            for i in range(7)
        ]
        assert min(widths) >= 0.45, f"слишком узкая колонка: {widths}"
        # sum(column widths) ~ landscape usable width (10.5")
        assert 9.8 <= sum(widths) <= 11.2, widths
        # header row: bold + повтор на страницах (tblHeader)
        assert findings.rows[0].cells[0].paragraphs[0].runs[0].bold is True
        trPr = findings.rows[0]._tr.trPr
        assert trPr is not None and trPr.find(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblHeader"
        ) is not None
        # small font for dense table
        run = findings.rows[1].cells[5].paragraphs[0].runs[0]
        assert run.font.size.pt <= 8.0
        # URL переносится zero-width-space'ами, но читается без них
        all_urls = [r.cells[5].text.replace("\u200b", "") for r in findings.rows[1:]]
        assert any("x" * 120 in u for u in all_urls), all_urls
        # evidence присутствует
        all_ev = [r.cells[6].text.replace("\u200b", "") for r in findings.rows[1:]]
        assert any("SQL error" in e for e in all_ev)
    finally:
        for f in (tmp_db, out):
            if os.path.exists(f):
                os.remove(f)


def test_docx_empty_db_does_not_crash():
    tmp_db = tempfile.mktemp(suffix=".db")
    out = tempfile.mktemp(suffix=".docx")
    try:
        db.set_db_path(tmp_db)
        db.init_db()
        n = report_export.export_docx(out)
        assert n == 0
        from docx import Document
        assert Document(out).tables
    finally:
        for f in (tmp_db, out):
            if os.path.exists(f):
                os.remove(f)


def test_html_report_title_typo_fixed():
    tmp_db = tempfile.mktemp(suffix=".db")
    out = tempfile.mktemp(suffix=".html")
    try:
        db.set_db_path(tmp_db)
        db.init_db()
        db.insert_scan_result(
            template_id="t1", worker_id=0, method="GET",
            url="http://target.test/x", status_code=200, error=None,
            check_type="xss_reflected", payload="<script>", evidence="reflected",
        )
        report_export.export_html(out)
        html = open(out, encoding="utf-8").read()
        assert "Notscan Scan Report" in html
        assert "Notsca Scan Report" not in html
    finally:
        for f in (tmp_db, out):
            if os.path.exists(f):
                os.remove(f)

def test_docx_report_is_localized_in_russian():
    import i18n
    from docx import Document

    tmp_db = tempfile.mktemp(suffix=".db")
    out = tempfile.mktemp(suffix=".docx")
    saved = i18n.get_language()
    try:
        _seed(tmp_db)
        i18n.set_language("ru")
        report_export.export_docx(out)
        text = "\n".join(p.text for p in Document(out).paragraphs)
        assert "Отчёт сканирования Notscan" in text
        assert "Сводка" in text
        assert "Находки" in text
        tables_text = "\n".join(
            c.text for t in Document(out).tables for row in t.rows for c in row.cells
        )
        assert "Критичность" in tables_text
        assert "Рекомендация" in tables_text
    finally:
        i18n.set_language(saved)
        for f in (tmp_db, out):
            if os.path.exists(f):
                os.remove(f)
