"""
report_export.py — Экспорт результатов скана в JSON, Markdown, DOCX, HTML и SARIF.

MD / JSON / DOCX содержат:
  • сводку по severity;
  • рекомендации «на что обратить внимание» (по типам находок);
  • матрицу детекции True/False (какой тип проверки сработал, а какой нет).
"""

from __future__ import annotations

import json
from datetime import datetime, timezone
from pathlib import Path
from typing import Optional

import db


def _now_iso() -> str:
    return datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")


def _catalog_severity_map() -> dict:
    """check_id → severity из единого каталога чеков (если доступен)."""
    try:
        from checks.catalog import CHECK_CATALOG  # локальный импорт: каталог лёгкий
        return {d.check_id.lower(): str(d.severity).upper() for d in CHECK_CATALOG}
    except Exception:
        return {}


_CATALOG_SEVERITY = _catalog_severity_map()


def _row_severity(r: dict) -> str:
    """Severity из scan_meta (CRITICAL/HIGH/MEDIUM/LOW/INFO) с фолбэком по check_type."""
    sev = (r.get("severity") or "").strip().upper()
    if sev:
        return sev
    ct = (r.get("check_type") or "").lower()
    if ct in _CATALOG_SEVERITY:
        return _CATALOG_SEVERITY[ct]
    if ct in ("cmdi", "ssti", "deserialization", "file_upload",
              "sqli", "sqli_time", "sqli_blind", "ssi_injection", "xxe", "ssrf",
              "path_traversal"):
        return "HIGH"
    if ct in ("xss", "xss_reflected", "xss_dom", "open_redirect", "crlf_injection",
              "html_injection", "nosql_injection", "ldap_injection",
              "xpath_injection", "idor", "cors", "ria_policy", "http_methods"):
        return "MEDIUM"
    if ct in ("csrf", "clickjacking", "security_headers",
              "info_disclosure", "outdated_libs", "third_party_scripts", "metafiles",
              "cookies", "hsts", "csp", "cache_control"):
        return "LOW"
    return "INFO"


def _sarif_level(sev: str) -> str:
    if sev in ("CRITICAL", "HIGH"):
        return "error"
    if sev in ("MEDIUM", "LOW"):
        return "warning"
    return "note"


# ──────────────────────────── Рекомендации ────────────────────────────

RECOMMENDATIONS = {
    "sqli": "Параметризованные запросы (prepared statements). Не конкатенируй пользовательский ввод в SQL.",
    "sqli_time": "Параметризованные запросы + таймауты БД. Валидация и строгая типизация ввода.",
    "cmdi": "Не передавай пользовательский ввод в shell/exec. Allowlist команд и разделение аргументов.",
    "ssti": "Не используй eval пользовательского ввода в шаблонах. Sandbox и allowlist выражений.",
    "xss": "Контекстное экранирование вывода (HTML/JS/attribute). CSP и HttpOnly cookies.",
    "xss_reflected": "Экранируй отражённый вывод + CSP. Не выводи пользовательский ввод без фильтрации.",
    "xss_dom": "Безопасные DOM-API (textContent вместо innerHTML). Избегай sinks: eval, document.write, innerHTML.",
    "path_traversal": "Канонизация путей + allowlist. Не принимай пути напрямую от клиента.",
    "xxe": "Отключи обработку внешних сущностей (DTD/ENTITY) в XML-парсере.",
    "ssrf": "Allowlist исходящих адресов. Блокируй metadata-эндпоинты (169.254.169.254). Валидируй URL.",
    "nosql_injection": "Валидация и типизация ввода для NoSQL. Не подставляй объекты в запросы напрямую.",
    "ldap_injection": "Экранируй спецсимволы LDAP или используй безопасные API с фильтрами.",
    "xpath_injection": "Параметризованные XPath-запросы (переменные, а не конкатенация).",
    "deserialization": "Не десериализуй недоверенные данные. Allowlist классов + подписи.",
    "crlf_injection": "Санитизируй CR/LF в заголовках. Не вставляй пользовательский ввод в заголовки ответа.",
    "open_redirect": "Allowlist URL для редиректов. Не доверяй параметрам next/return/url без проверки.",
    "file_upload": "Валидация типа по содержимому (не по расширению). Храни загрузки вне webroot.",
    "idor": "Проверка прав на уровне объекта (object-level authorization). Не доверяй ID из URL.",
    "jwt": "Валидируй подпись JWT на сервере. Отклоняй alg=none. Используй длинный случайный HMAC-секрет или RS256.",
    "info_disclosure_paths": "Закрой доступ к чувствительным путям (actuator, .env, .git, swagger, metrics и т.п.) или ограничь их авторизацией.",
    "sqli_blind": "Параметризованные запросы (prepared statements). Не конкатенируй пользовательский ввод в SQL; следи за одинаковым поведением ответа на булевы условия.",
    "ssi_injection": "Не обрабатывай пользовательский ввод SSI-движком; экранируй директивы или отключи SSI для пользовательских данных.",
    "html_injection": "Контекстное HTML-экранирование вывода. Не вставляй пользовательский ввод в HTML без кодирования.",
    "cookies": "Установи Secure, HttpOnly и SameSite=Lax/Strict для сессионных cookie; не храни секреты в cookie.",
    "hsts": "Добавь Strict-Transport-Security (max-age >= 31536000, includeSubDomains) для HTTPS-ответов.",
    "csp": "Настрой Content-Security-Policy без unsafe-inline/unsafe-eval; используй nonce/hash и строгий script-src.",
    "cache_control": "Для ответов с сессией/секретами используй Cache-Control: no-store, private и Pragma: no-cache.",
    "insecure_login": "Передавай креды только по HTTPS; не используй формы логина и Authorization на HTTP.",
    "browser_storage": "Не храни токены/секреты в localStorage/sessionStorage; используй HttpOnly-cookie с защитой от XSS.",
    "metafiles": "Проверь robots.txt/sitemap.xml/security.txt на утечку внутренних путей и удали лишние записи.",
    "ria_policy": "Убери wildcard из crossdomain.xml/clientaccesspolicy.xml; ограничь домены явным allowlist.",
    "http_methods": "Отключи ненужные HTTP-методы (TRACE/PUT/DELETE), оставь только необходимые; ограничь через прокси/сервер.",
    "sensitive_files": "Закрой доступ к бэкапам, .env/.git и служебным файлам; храни их вне webroot.",
    "cors": "Не отражай произвольный Origin и не используй Access-Control-Allow-Credentials вместе с wildcard; веди allowlist origin.",
    "csrf": "CSRF-токены + SameSite cookies. Проверяй Origin/Referer.",
    "clickjacking": "Заголовки X-Frame-Options / CSP frame-ancestors.",
    "security_headers": "Добавь недостающие security-заголовки (HSTS, X-Content-Type-Options, CSP и др.).",
    "info_disclosure": "Скрывай версии, внутренние пути и stack traces в продакшене.",
    "outdated_libs": "Обнови устаревшие JS-библиотеки/фреймворки.",
    "third_party_scripts": "Ограничь сторонние скрипты (SRI, allowlist доменов).",
}

GENERIC_RECOMMENDATION = "Проверь вручную и устрани в соответствии с OWASP Testing Guide / CWE."


def _recommendation_for(check_type: str) -> str:
    return RECOMMENDATIONS.get((check_type or "").lower(), GENERIC_RECOMMENDATION)


# ──────────────────────────── Данные отчёта ────────────────────────────

_SEVERITY_ORDER = ["CRITICAL", "HIGH", "MEDIUM", "LOW", "INFO"]


def _sev_rank(sev: str) -> int:
    try:
        return _SEVERITY_ORDER.index(sev)
    except ValueError:
        return len(_SEVERITY_ORDER)


def _dedup_findings(findings: list[dict]) -> list[dict]:
    """
    Сворачивает дубликаты находок по стабильному fingerprint (или fallback
    check_type|url), повторяя поведение Results-вью «Group duplicates».
    Оставляет представителя группы с максимальной severity и добавляет
    поле `occurrences` (сколько находок представляет данная запись).
    """
    reps: dict[str, dict] = {}
    order: list[str] = []
    counts: dict[str, int] = {}
    for r in findings:
        fp = r.get("fingerprint")
        key = fp if fp else f"{r.get('check_type')}|{r.get('url')}"
        if key not in reps:
            reps[key] = r
            order.append(key)
            counts[key] = 1
        else:
            counts[key] += 1
            if _sev_rank(r.get("severity")) < _sev_rank(reps[key].get("severity")):
                reps[key] = r
    out: list[dict] = []
    for k in order:
        r = dict(reps[k])
        r["occurrences"] = counts[k]
        out.append(r)
    return out


def _enriched_findings(rows: list[dict]) -> list[dict]:
    """Enrich evidence-bearing rows with severity + recommendation."""
    findings = [r for r in rows if r.get("evidence")]
    return [
        {
            **r,
            "severity": _row_severity(r),
            "recommendation": _recommendation_for(r.get("check_type")),
        }
        for r in findings
    ]


def _build_report_data(limit: int | None = None, deduplicate: bool = True) -> dict:
    rows = db.fetch_scan_results_with_meta(limit=limit)
    findings_raw = _enriched_findings(rows)
    findings = _dedup_findings(findings_raw) if deduplicate else findings_raw

    sev_counts = {s: 0 for s in _SEVERITY_ORDER}
    by_type: dict[str, dict] = {}

    for r in findings:
        sev = _row_severity(r)
        sev_counts[sev] = sev_counts.get(sev, 0) + 1
        ct = (r.get("check_type") or "unknown").lower()
        d = by_type.setdefault(ct, {"count": 0, "max_severity": "INFO", "rank": len(_SEVERITY_ORDER)})
        d["count"] += 1
        if _sev_rank(sev) < d["rank"]:
            d["rank"] = _sev_rank(sev)
            d["max_severity"] = sev

    recommendations = sorted(
        [
            {
                "check_type": ct,
                "count": d["count"],
                "max_severity": d["max_severity"],
                "recommendation": _recommendation_for(ct),
            }
            for ct, d in by_type.items()
        ],
        key=lambda x: (-x["count"], _sev_rank(x["max_severity"])),
    )

    all_types = sorted({(r.get("check_type") or "unknown").lower() for r in rows})
    matrix = []
    for ct in all_types:
        d = by_type.get(ct)
        matrix.append({
            "check_type": ct,
            "detected": ct in by_type,
            "count": d["count"] if d else 0,
            "max_severity": d["max_severity"] if d else "",
        })

    return {
        "total": len(rows),
        "findings_count": len(findings),
        "findings_raw_count": len(findings_raw),
        "deduplicated": deduplicate,
        "severity_counts": sev_counts,
        "findings": findings,
        "recommendations": recommendations,
        "matrix": matrix,
    }


def _md_escape(s) -> str:
    return (s or "").replace("|", "\\|").replace("\n", " ").replace("\r", " ")


# ──────────────────────────── JSON ────────────────────────────

def export_json(filepath: str | Path, limit: int | None = None, deduplicate: bool = True) -> int:
    data = _build_report_data(limit, deduplicate)
    out = {
        "export_time": _now_iso(),
        "source": "notsca",
        "summary": {
            "total": data["total"],
            "findings": data["findings_count"],
            "findings_raw": data["findings_raw_count"],
            "deduplicated": data["deduplicated"],
            "severity_counts": data["severity_counts"],
        },
        "recommendations": data["recommendations"],
        "detection_matrix": data["matrix"],
        "results": [
            {
                "id": f["id"],
                "time": f["time"],
                "method": f["method"],
                "url": f["url"],
                "status_code": f["status_code"],
                "check_type": f["check_type"],
                "payload": f["payload"],
                "evidence": f["evidence"],
                "severity": f["severity"],
                "confidence": f.get("confidence"),
                "occurrences": f.get("occurrences", 1),
                "recommendation": f["recommendation"],
            }
            for f in data["findings"]
        ],
    }
    path = Path(filepath)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")
    return len(data["findings"])


# ──────────────────────────── Markdown ────────────────────────────

def export_md(filepath: str | Path, limit: int | None = None, deduplicate: bool = True) -> int:
    data = _build_report_data(limit, deduplicate)
    L: list[str] = []
    L.append("# Notscan Scan Report")
    L.append("")
    dedup_note = (
        f"  |  Unique findings: {data['findings_count']} (raw: {data['findings_raw_count']})"
        if data.get("deduplicated") else ""
    )
    L.append(
        f"> Generated: {_now_iso()}  |  Total results: {data['total']}  |  Findings: {data['findings_count']}"
        f"{dedup_note}"
    )
    L.append("")

    L.append("## Summary")
    L.append("")
    L.append("| Severity | Count |")
    L.append("|----------|-------|")
    for sev in _SEVERITY_ORDER:
        L.append(f"| {sev.title()} | {data['severity_counts'].get(sev, 0)} |")
    L.append("")

    L.append("## Recommendations (на что обратить внимание)")
    L.append("")
    L.append("| Check | Severity | Count | Recommendation |")
    L.append("|-------|----------|-------|----------------|")
    for rec in data["recommendations"]:
        L.append(
            f"| {rec['check_type']} | {rec['max_severity'].title()} | {rec['count']} | {rec['recommendation']} |"
        )
    L.append("")

    L.append("## Detection Matrix (True/False)")
    L.append("")
    L.append("| Check | Detected | Count | Max Severity |")
    L.append("|-------|----------|-------|--------------|")
    for m in data["matrix"]:
        det = "True" if m["detected"] else "False"
        mx = m["max_severity"].title() if m["max_severity"] else "—"
        L.append(f"| {m['check_type']} | {det} | {m['count']} | {mx} |")
    L.append("")

    L.append("## Findings")
    L.append("")
    L.append("| ID | Check | Severity | Method | Count | URL | Evidence |")
    L.append("|----|-------|----------|--------|-------|-----|----------|")
    for f in data["findings"]:
        cnt = f.get("occurrences", 1)
        L.append(
            f"| {f['id']} | {f.get('check_type')} | {f['severity'].title()} | "
            f"{f.get('method')} | {cnt} | {_md_escape(f.get('url', ''))} | "
            f"{_md_escape((f.get('evidence') or '')[:200])} |"
        )
    L.append("")

    Path(filepath).write_text("\n".join(L), encoding="utf-8")
    return len(data["findings"])


# ──────────────────────────── DOCX ────────────────────────────

def _docx_table(doc, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    try:
        table.style = "Light Grid Accent 1"
    except Exception:
        pass
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for row in rows:
        cells = table.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)


def export_docx(filepath: str | Path, limit: int | None = None) -> int:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("python-docx не установлен. Выполни: pip install python-docx") from exc

    data = _build_report_data(limit)
    doc = Document()
    doc.add_heading("Notscan Scan Report", 0)
    dedup_txt = (
        f"   |   Unique findings: {data['findings_count']} (raw: {data['findings_raw_count']})"
        if data.get("deduplicated") else ""
    )
    doc.add_paragraph(
        f"Generated: {_now_iso()}   |   Total results: {data['total']}   |   Findings: {data['findings_count']}"
        f"{dedup_txt}"
    )

    doc.add_heading("Summary", level=1)
    _docx_table(
        doc,
        ["Severity", "Count"],
        [[sev.title(), str(data["severity_counts"].get(sev, 0))] for sev in _SEVERITY_ORDER],
    )

    doc.add_heading("Recommendations", level=1)
    _docx_table(
        doc,
        ["Check", "Severity", "Count", "Recommendation"],
        [
            [rec["check_type"], rec["max_severity"].title(), str(rec["count"]), rec["recommendation"]]
            for rec in data["recommendations"]
        ],
    )

    doc.add_heading("Detection Matrix (True/False)", level=1)
    _docx_table(
        doc,
        ["Check", "Detected", "Count", "Max Severity"],
        [
            [m["check_type"], "True" if m["detected"] else "False", str(m["count"]),
             m["max_severity"].title() if m["max_severity"] else "—"]
            for m in data["matrix"]
        ],
    )

    doc.add_heading("Findings", level=1)
    _docx_table(
        doc,
        ["ID", "Check", "Severity", "Method", "Count", "URL", "Evidence"],
        [
            [str(f["id"]), str(f.get("check_type")), f["severity"].title(),
             str(f.get("method")), str(f.get("occurrences", 1)),
             str(f.get("url", "")), (f.get("evidence") or "")[:300]]
            for f in data["findings"]
        ],
    )

    doc.save(filepath)
    return len(data["findings"])


# ──────────────────────────── HTML ────────────────────────────

def export_html(filepath: str | Path, limit: int | None = None, deduplicate: bool = True) -> int:
    rows = db.fetch_scan_results_with_meta(limit=limit)
    path = Path(filepath)
    path.parent.mkdir(parents=True, exist_ok=True)

    findings = _enriched_findings(rows)
    raw_count = len(findings)
    vuln_rows = _dedup_findings(findings) if deduplicate else findings
    html_parts = [
        "<!DOCTYPE html><html><head><meta charset='utf-8'>",
        "<title>Notsca Scan Report</title>",
        "<style>"
        "body{font-family:Segoe UI,sans-serif;background:#1e1e2e;color:#cdd6f4;margin:20px;}",
        "h1{color:#89b4fa;} .meta{color:#a6adc8;margin-bottom:20px;}",
        "table{border-collapse:collapse;width:100%;margin-bottom:20px;}",
        "th,td{border:1px solid #45475a;padding:8px;text-align:left;}",
        "th{background:#313244;color:#89b4fa;}",
        ".sev-error{color:#f38ba8;} .sev-warning{color:#f9e2af;} .sev-note{color:#a6adc8;}",
        "pre{overflow-x:auto;font-size:12px;}",
        "</style></head><body>",
        "<h1>Notsca Scan Report</h1>",
        f"<p class='meta'>Export: {_now_iso()} | Total: {len(rows)} | "
        f"Unique findings: {len(vuln_rows)} (raw: {raw_count})</p>",
        "<table><thead><tr>",
        "<th>ID</th><th>Time</th><th>Check</th><th>Severity</th><th>Method</th><th>Count</th><th>URL</th>",
        "<th>Payload</th><th>Evidence</th>",
        "</tr></thead><tbody>",
    ]
    for r in vuln_rows:
        sev = _row_severity(r)
        level = _sarif_level(sev)
        ts = datetime.fromtimestamp(r["time"]).strftime("%Y-%m-%d %H:%M:%S") if r.get("time") else ""
        html_parts.append(
            f"<tr class='sev-{level}'>"
            f"<td>{r['id']}</td><td>{ts}</td><td>{r.get('check_type', '')}</td>"
            f"<td class='sev-{level}'>{sev}</td><td>{r.get('method', '')}</td>"
            f"<td>{r.get('occurrences', 1)}</td><td>{r.get('url', '')}</td>"
            f"<td><pre>{json.dumps(r.get('payload') or '', ensure_ascii=False)[:200]}</pre></td>"
            f"<td><pre>{json.dumps(r.get('evidence') or '', ensure_ascii=False)[:500]}</pre></td>"
            "</tr>"
        )
    html_parts.append("</tbody></table></body></html>")
    path.write_text("".join(html_parts), encoding="utf-8")
    return len(vuln_rows)


# ──────────────────────────── SARIF ────────────────────────────

def export_sarif(filepath: str | Path, limit: int | None = None, deduplicate: bool = True) -> int:
    rows = db.fetch_scan_results_with_meta(limit=limit)
    findings = _enriched_findings(rows)
    vuln_rows = _dedup_findings(findings) if deduplicate else findings
    path = Path(filepath)
    path.parent.mkdir(parents=True, exist_ok=True)

    results_sarif = []
    for r in vuln_rows:
        sev = _row_severity(r)
        rule_id = (r.get("check_type") or "unknown").replace(" ", "_")
        occ = r.get("occurrences", 1)
        msg = (r.get("evidence") or "")[:1000]
        if occ > 1:
            msg = f"[×{occ}] {msg}"
        results_sarif.append({
            "ruleId": rule_id,
            "level": _sarif_level(sev),
            "message": {"text": msg},
            "locations": [{
                "physicalLocation": {
                    "artifactLocation": {"uri": r.get("url") or ""},
                    "region": {"startLine": 1},
                }
            }],
            "partialFingerprints": {
                "notsca/result": r.get("fingerprint") or str(r.get("id", "")),
            },
        })

    rules_seen = {r.get("check_type") or "unknown" for r in vuln_rows}
    rules = [
        {"id": rid, "name": rid, "shortDescription": {"text": rid}}
        for rid in rules_seen
    ]

    sarif = {
        "$schema": "https://raw.githubusercontent.com/oasis-tcs/sarif-spec/master/Schemata/sarif-schema-2.1.0.json",
        "version": "2.1.0",
        "runs": [{
            "tool": {
                "driver": {
                    "name": "notsca",
                    "version": "1.0",
                    "informationUri": "https://github.com/notsca/notsca",
                    "rules": rules,
                }
            },
            "results": results_sarif,
        }]
    }
    path.write_text(json.dumps(sarif, ensure_ascii=False, indent=2), encoding="utf-8")
    return len(vuln_rows)
